from copy import deepcopy
from pathlib import Path
import re
from lxml import etree
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\Admin\Documents\ChatGPT\HTQLHDK")
BASE = ROOT / "outputs" / "BaoCaoPT&TKHTTT_Nhom20_28092026.docx"
EXPANDED = ROOT / "outputs" / "BaoCaoPT_TKHTTT_Nhom20_28092026_ban_viet_lai.docx"
OUT = ROOT / "outputs" / "BaoCaoPT_TKHTTT_Nhom20_theo_mau_khong_trich_dan.docx"

base = Document(BASE)
source = Document(EXPANDED)
body = base._element.body
source_body = source._element.body

def find_heading(doc, prefix):
    for p in doc.paragraphs:
        if p.style.name == "Heading 1" and p.text.startswith(prefix):
            return p._p
    raise ValueError(f"Không tìm thấy đề mục: {prefix}")

old_ch1 = find_heading(base, "CHƯƠNG I.")
old_ch4 = find_heading(base, "CHƯƠNG IV.")
new_ch1 = find_heading(source, "CHƯƠNG I.")
new_refs = find_heading(source, "TÀI LIỆU THAM KHẢO")

def body_slice(start, end, the_body):
    children = list(the_body)
    return children[children.index(start):children.index(end)]

def after_ch4_signature(doc):
    ch4 = find_heading(doc, "CHƯƠNG IV.")
    elems = list(doc._element.body)
    return b"".join(etree.tostring(x, method="c14n") for x in elems[elems.index(ch4):-1])

preserved_tail = after_ch4_signature(base)

# Xóa đúng phần Chương I–III của tệp Nhóm 20. Tất cả phần tử từ Chương IV
# tới Tài liệu tham khảo (gồm Chương VI và bảng mẫu) nằm ngoài vùng thay thế.
for element in body_slice(old_ch1, old_ch4, body):
    body.remove(element)

figure_files = [
    ("architecture/rendered/02_container_components.png", 15.5),
    ("so_do_cap_nhat/rendered/02_activity_quy_trinh_tong_the.png", 13.2),
    ("visual_paradigm_uml_md/rendered/18_use_case_tai_san_lich_quet.png", 8.2),
    ("visual_paradigm_uml_md/rendered/19_use_case_nhap_bao_cao_zap.png", 11.3),
    ("visual_paradigm_uml_md/rendered/21_use_case_khac_phuc_xac_minh.png", 9.2),
    ("visual_paradigm_uml_md/rendered/03_activity_khai_bao_phe_duyet_tai_san.png", 10.8),
    ("visual_paradigm_uml_md/rendered/07_activity_khac_phuc_va_xac_minh.png", 10.7),
    ("so_do_cap_nhat/rendered/03_state_zap_execution.png", 15.5),
    ("so_do_cap_nhat/rendered/04_dfd_level_1.png", 15.5),
    ("_report_work/erd_13_bang_bao_cao.png", 15.7),
]

def clean_text(value):
    value = re.sub(r"\s*\((?:nguồn|biên soạn từ):?\s*\[[^)]*\]\)", "", value, flags=re.I)
    value = re.sub(r"\s*\[(?:[1-8])(?:\s*,\s*[1-8])*\]", "", value)
    value = re.sub(r"\s+([.,;:])", r"\1", value)
    value = re.sub(r"\.\.$", ".", value)
    return value

source_parts = body_slice(new_ch1, new_refs, source_body)
image_index = 0
inserted_count = 0
for part in source_parts:
    if part.tag == qn("w:p") and part.xpath(".//w:drawing"):
        rel, width = figure_files[image_index]
        image_index += 1
        base.add_picture(str(ROOT / rel), width=Cm(width))
        figure_paragraph = base.paragraphs[-1]
        figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure_paragraph.paragraph_format.keep_with_next = True
        p = figure_paragraph._p
        old_ch4.addprevious(p)
    else:
        clone = deepcopy(part)
        for t in clone.xpath(".//w:t"):
            if t.text: t.text = clean_text(t.text)
        old_ch4.addprevious(clone)
    inserted_count += 1

assert image_index == 10, image_index
assert after_ch4_signature(base) == preserved_tail, "Phần Chương IV trở đi đã thay đổi trong bộ nhớ"

# Cập nhật mục lục chữ của Chương I–III. Chương IV trở đi chỉ được nhắc đến,
# không sửa chính văn của các chương đó.
abbr = find_heading(base, "DANH MỤC TỪ VIẾT TẮT")
toc = find_heading(base, "MỤC LỤC")
for element in body_slice(toc, abbr, body)[1:]:
    body.remove(element)
toc_lines = [
    "CHƯƠNG I. TỔNG QUAN",
    "    1.1. Bối cảnh và lý do chọn đề tài",
    "    1.2. Mục tiêu và kết quả kỳ vọng",
    "    1.3. Phạm vi và giới hạn",
    "    1.4. Cơ sở lý thuyết",
    "    1.5. Nền tảng kỹ thuật và hiện trạng",
    "CHƯƠNG II. PHÂN TÍCH NỘI DUNG, YÊU CẦU",
    "    2.1. Tác nhân, quyền và trách nhiệm",
    "    2.2–2.6. Quy trình nghiệp vụ F-01 đến F-04",
    "    2.7–2.8. Yêu cầu chức năng và chất lượng",
    "CHƯƠNG III. PHÂN TÍCH THIẾT KẾ HỆ THỐNG",
    "    3.1–3.4. Use case, hoạt động, trạng thái và kiến trúc",
    "    3.5–3.7. ERD, danh sách bảng và đối chiếu dữ liệu",
    "    3.8–3.10. Lớp miền, giao diện và kiểm chứng",
    "CHƯƠNG IV. PHÁT TRIỂN/THỰC THI",
    "CHƯƠNG V. TRIỂN KHAI",
    "CHƯƠNG VI. KẾT LUẬN",
    "TÀI LIỆU THAM KHẢO",
]
for line in toc_lines:
    p = base.add_paragraph()
    p.add_run(line).bold = line.startswith("CHƯƠNG") or line.startswith("TÀI LIỆU")
    abbr.addprevious(p._p)

# Danh sách hình/bảng là thành phần của mẫu. Không kèm ghi nguồn.
figure_labels = []
table_labels = []
in_chapters = False
for p in base.paragraphs:
    if p.text.startswith("CHƯƠNG I."): in_chapters = True
    if p.text.startswith("CHƯƠNG IV."): in_chapters = False
    if in_chapters:
        if p.text.startswith("Hình "): figure_labels.append(clean_text(p.text))
        if p.text.startswith("Bảng "): table_labels.append(clean_text(p.text))

heading = base.add_paragraph(style="Heading 1")
heading.add_run("DANH SÁCH HÌNH, BẢNG")
abbr.addprevious(heading._p)
for label in figure_labels + table_labels:
    p = base.add_paragraph()
    p.add_run(label)
    abbr.addprevious(p._p)

assert after_ch4_signature(base) == preserved_tail, "Phần Chương IV trở đi đã thay đổi"
base.save(OUT)
print(OUT)
print(f"Đã thay {inserted_count} phần tử Chương I–III; {image_index} hình; {len(figure_labels)} nhãn hình; {len(table_labels)} nhãn bảng")
