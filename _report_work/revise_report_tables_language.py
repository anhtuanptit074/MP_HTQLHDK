from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\Admin\Documents\ChatGPT\HTQLHDK")
SOURCE = ROOT / "outputs" / "BaoCaoPT_TKHTTT_Nhom20_theo_mau_khong_trich_dan.docx"
OUT = ROOT / "outputs" / "BaoCaoPT_TKHTTT_Nhom20_chu_bang_12_tieng_viet.docx"
doc = Document(SOURCE)
body = doc._element.body

def heading(prefix, style="Heading 1"):
    for p in doc.paragraphs:
        if p.style.name == style and p.text.startswith(prefix): return p
    raise ValueError(prefix)

ch1 = heading("CHƯƠNG I.")
ch4 = heading("CHƯƠNG IV.")

# Xóa toàn bộ mục đối chiếu trong Chương III, gồm bảng và đoạn kết luận đi kèm.
compare = heading("3.7. Đối chiếu", "Heading 2")
next_section = heading("3.8. Sơ đồ lớp", "Heading 2")
elements = list(body)
for element in elements[elements.index(compare._p):elements.index(next_section._p)]:
    body.remove(element)

def remove_last_column(table):
    grid = table._tbl.tblGrid
    grid.remove(list(grid)[-1])
    for row in table.rows:
        tc = row.cells[-1]._tc
        row._tr.remove(tc)

uc = next(t for t in doc.tables if len(t.rows) == 64 and t.rows[0].cells[-1].text == "Phạm vi")
ad = next(t for t in doc.tables if len(t.rows) == 17 and t.rows[0].cells[-1].text == "Phạm vi")
remove_last_column(uc)
remove_last_column(ad)
for table, widths in [(uc,[2.0,6.3,7.5]),(ad,[2.1,13.7])]:
    for i,width in enumerate(widths): table.columns[i].width = Cm(width)
    for row in table.rows:
        for i,width in enumerate(widths): row.cells[i].width = Cm(width)

def paragraph_in_chapters(p):
    elems = list(body)
    return elems.index(ch1._p) <= elems.index(p._p) < elems.index(ch4._p)

replacements = [
    (r"\buse cases?\b", "ca sử dụng"),
    (r"\bactivity diagrams?\b", "sơ đồ hoạt động"),
    (r"\bbaseline\b", "thiết kế ban đầu"),
    (r"\bWorker\b", "tiến trình nền"),
    (r"\bscope\b", "phạm vi"),
    (r"\bscan window\b", "khung giờ quét"),
    (r"\bprofile\b", "hồ sơ quét"),
    (r"\bartifact\b", "tệp báo cáo"),
    (r"\bparser\b", "bộ phân tích báo cáo"),
    (r"\bscanner\b", "công cụ quét"),
    (r"\bswimlane\b", "làn trách nhiệm"),
    (r"\bguard\b", "điều kiện kiểm tra"),
    (r"\bdatabase\b", "cơ sở dữ liệu"),
    (r"\bfrontend\b", "giao diện người dùng"),
    (r"\bbackend\b", "máy chủ ứng dụng"),
    (r"\bend-to-end\b", "đầu cuối"),
    (r"\bproduction\b", "môi trường vận hành thật"),
    (r"\bstaging\b", "môi trường thử nghiệm"),
    (r"\brunner\b", "bộ thực thi"),
    (r"\bobject storage\b", "kho tệp"),
    (r"\bsecret\b", "bí mật cấu hình"),
    (r"\bworkflow\b", "quy trình xử lý"),
    (r"\brollback\b", "hoàn tác"),
    (r"\btransaction\b", "giao dịch"),
    (r"\bschema\b", "lược đồ dữ liệu"),
    (r"\bscan run\b", "lần quét"),
    (r"\bfinding\b", "phát hiện"),
    (r"\bticket\b", "phiếu"),
    (r"\baudit\b", "nhật ký kiểm toán"),
    (r"\bplan\b", "kế hoạch"),
    (r"\brun\b", "lần chạy"),
]

def vi(text):
    before = text
    for pattern,replacement in replacements:
        text = re.sub(pattern,replacement,text,flags=re.I)
    text = re.sub(r",\s*,\s*\.", ".", text)
    text = re.sub(r",\s*\.", ".", text)
    text = re.sub(r",\s*,", ",", text)
    text = re.sub(r"\s+([.,;:])", r"\1", text)
    if before and before[0].isupper() and text and text[0].islower():
        text = text[0].upper()+text[1:]
    return text

# Sửa ba đoạn còn thiên về so sánh; phần còn lại dùng câu chữ tiếng Việt.
specific = {
    "Phần cập nhật ngày 28/09/2026": "Bản cập nhật ngày 28/09/2026 bổ sung cơ chế điều phối OWASP ZAP bằng tiến trình nền và mô hình dữ liệu đích 13 bảng. Báo cáo tập trung mô tả mục tiêu, quy trình và thiết kế của hệ thống theo bốn chức năng chính.",
    "V002 tách scan_occurrence": "Trong phương án quét tự động, hệ thống lưu trạng thái nghiệp vụ của lần quét và trạng thái kỹ thuật của phiên ZAP. Khi báo cáo hợp lệ, tiến trình nền lưu tệp kết quả, cập nhật lần quét và tạo lô nhập ở trạng thái Received trong cùng giao dịch. Bộ phân tích báo cáo F-03 tiếp tục xử lý lô nhập.",
    "Hợp đồng baseline 27 bảng": "Báo cáo ZAP phải được kiểm tra toàn bộ về định dạng, tài sản và phạm vi. Nếu dữ liệu nguồn không hợp lệ, hệ thống từ chối hoặc hoàn tác giao dịch và không tạo phát hiện. Báo cáo hợp lệ được chuẩn hóa và ghi nhận theo từng lần quét.",
}

final_rewrites = {
    "Một website có thể sinh": "Một trang web có thể sinh nhiều phát hiện bảo mật qua các kỳ quét. Nếu chỉ lưu tệp kết quả, đơn vị khó xác định phát hiện thuộc tài sản nào, phạm vi quét đã được chấp thuận hay chưa, ai chịu trách nhiệm và việc khắc phục đã được xác minh hay chưa. HTQLHDK quản lý toàn bộ vòng đời đó, từ khai báo tài sản đến theo dõi và đóng phát hiện.",
    "• Quản lý tài sản website": "• Quản lý tài sản trang web cùng phiên bản phạm vi, quy tắc cho phép hoặc loại trừ URL và khung giờ quét đã được phê duyệt.",
    "Một phát hiện bảo mật trải qua": "Một phát hiện bảo mật trải qua các bước ghi nhận, phân tích, phân công, khắc phục và xác minh. Mức rủi ro (Risk) và độ tin cậy (Confidence) do ZAP cung cấp là thông tin kỹ thuật; mức độ nghiêm trọng cuối cùng (Final Severity) là quyết định nghiệp vụ có lý do và người xác nhận. Dấu vân tay của phát hiện giúp chống trùng qua các kỳ quét. Việc không thấy phát hiện trong lần quét sau chỉ có giá trị xác minh khi vùng bao phủ quét đã được xác nhận.",
    "Thiết kế cơ sở dữ liệu dùng khóa": "Thiết kế cơ sở dữ liệu dùng khóa, ràng buộc duy nhất và quan hệ riêng cho dữ liệu nhiều giá trị. Mô hình đích hướng đến dạng chuẩn BCNF cho phần lớn bảng, không lưu lặp kỳ quét kế tiếp, thời điểm xuất hiện đầu/cuối hoặc số liệu tổng hợp có thể tính bằng truy vấn. Nhật ký kiểm toán chỉ cho phép ghi thêm để bảo lưu lịch sử thay đổi nhạy cảm.",
    "Trùng domain/URL": "Trùng tên miền hoặc URL với tài sản khác là tín hiệu cần rà soát, không tự động dẫn đến từ chối. Khi thay đổi phạm vi, hệ thống tạo phiên bản mới; lịch hiện hành chỉ dùng phiên bản đã được duyệt và còn hiệu lực.",
    "• Chỉ quét tài sản Active": "• Chỉ quét tài sản ở trạng thái Active trong phạm vi Approved và khung giờ hợp lệ. Quy tắc Exclude được ưu tiên khi URL khớp cả quy tắc Include và Exclude.",
    "• hồ sơ quét ZAP": "• Hồ sơ quét ZAP phải có phiên bản. Bản đã kích hoạt hoặc ngừng sử dụng không được sửa; chế độ quét chủ động chỉ bật sau khi mục tiêu và cấu hình được phê duyệt.",
    "• Báo cáo nguồn và metadata": "• Báo cáo gốc và siêu dữ liệu kiểm tra toàn vẹn phải bất biến; nội dung Attack/Evidence, mã truy cập và URL nhạy cảm không được đưa vào nhật ký thô.",
    "• API kiểm lại quyền": "• API phải kiểm tra lại quyền theo hành động, tài sản, nhiệm vụ và trạng thái. Việc giao diện ẩn nút thao tác không thay thế kiểm soát quyền ở máy chủ.",
    "• Các thao tác ghi cần idempotency": "• Thao tác ghi phải chống xử lý lặp, dùng giao dịch và ghi nhật ký kiểm toán. Khi mất quyền giữ công việc (lease), hệ thống không được tự động quét lại mục tiêu.",
    "Mục tiêu định lượng như SLA": "Mục tiêu định lượng như SLA, thời hạn lưu báo cáo, giới hạn tải lên và RPO/RTO chưa được phê duyệt; báo cáo không tự đặt con số. Kiểm thử chấp nhận phải bao gồm trường hợp thành công, bị từ chối, hoàn tác, sai quyền và xử lý đồng thời.",
    "Sơ đồ ca sử dụng phân vai": "Sơ đồ ca sử dụng phân vai người dùng theo trách nhiệm. Những ca kiểm tra định dạng, chống trùng, che dữ liệu nhạy cảm và ghi nhật ký kiểm toán là xử lý nội bộ do ca sử dụng cha gọi, không phải thao tác độc lập. Bộ sơ đồ gốc có 63 ca sử dụng; nhóm chấp nhận rủi ro và một số chức năng báo cáo, thông báo thuộc phạm vi mở rộng so với ERD đích 13 bảng.",
    "Trong mô hình đích, ca sử dụng": "Trong mô hình đích, các ca quản lý tài sản, lập kế hoạch quét, nhận báo cáo và khắc phục được giữ lại. Chấp nhận rủi ro, chính sách SLA, hộp thông báo và liên kết chia sẻ báo cáo không có bảng riêng; nếu trở thành yêu cầu bắt buộc, cần bổ sung mô hình dữ liệu và luồng phê duyệt tương ứng.",
    "Sơ đồ hoạt động ở Hình 2.1": "Sơ đồ hoạt động ở Hình 2.1 bao quát luồng chính. Các sơ đồ AD-02, AD-03, AD-04 và AD-06 chỉ rõ trách nhiệm từng làn, điều kiện kiểm tra và nhánh lỗi. Phần quét tự động bổ sung tiến trình nền và ZAP vào luồng thực thi; quyết định nghiệp vụ của chủ tài sản và chuyên viên ATTT vẫn do con người thực hiện.",
    "Điểm quyết định quan trọng nhất": "Các điểm quyết định chính của AD-02 là quyền khai báo, URL hợp lệ và việc xác nhận, phê duyệt. AD-03 kiểm tra trạng thái tài sản, phạm vi đã duyệt và xung đột lịch. AD-04 kiểm tra quyền, khóa lô, định dạng và URL trước khi tạo phát hiện. AD-06 chỉ cho đóng khi kết quả xác minh gắn đúng tài sản, đúng phát hiện và đủ vùng bao phủ; kết quả không hợp lệ hoặc dấu vân tay còn xuất hiện sẽ dẫn tới mở lại.",
    "Ở lược đồ dữ liệu hiện hành": "Ở lược đồ dữ liệu hiện hành, tiến trình nền chuyển zap_execution qua các trạng thái kỹ thuật Queued → Preparing Plan → Starting Runner → Running → Collecting Report → Report Received. Failed, Timed Out và Cancelled là trạng thái kết thúc lỗi. scan_occurrence giữ trạng thái nghiệp vụ; tiến trình nền chỉ tạo import_batch ở trạng thái Received và không tự đóng phát hiện.",
    "Web UI gọi API": "Giao diện web gọi API để xác thực và kiểm quyền theo vai trò, tài sản và nhiệm vụ. PostgreSQL lưu trạng thái nghiệp vụ, quyền và nhật ký kiểm toán. Tiến trình nền riêng điều phối ZAP; công cụ quét không truy cập cơ sở dữ liệu nghiệp vụ. API không giữ Docker socket hoặc khóa ZAP. Môi trường vận hành thật cần bộ thực thi có quyền tối thiểu, kho tệp riêng tư và cơ chế quản lý khóa, bí mật cấu hình được phê duyệt.",
    "ERD dưới đây thể hiện": "ERD thể hiện các quan hệ cốt lõi của mô hình đích. app_user liên kết với user_role và các vai trò nghiệp vụ; asset có nhiều phiên bản phạm vi, mỗi phiên bản gồm quy tắc và khung giờ. scan_plan liên kết phạm vi với zap_scan_profile rồi sinh scan_run. finding_occurrence nối lần chạy với finding; remediation_ticket theo dõi khắc phục; audit_log lưu dấu vết. Kế hoạch quét loại Verification tham chiếu phiếu khắc phục.",
    "Mọi bảng vật lý trong mô hình đích": "Mỗi bảng vật lý trong mô hình đích dùng id làm khóa chính. Bảng dưới nêu khóa ứng viên, khóa ngoại quan trọng, thuộc tính chính và ràng buộc. Đây là thiết kế dữ liệu của đề tài, chưa phải khẳng định lược đồ SQL hiện hành đã chuyển đổi.",
    "Dữ liệu được suy ra thay vì lưu": "Các giá trị có thể tính bằng truy vấn gồm phạm vi hiện hành, kỳ quét kế tiếp, thời điểm xuất hiện đầu/cuối, số phát hiện theo mức độ, danh sách việc quá hạn và báo cáo tổng hợp. Các khung nhìn như v_asset_current_scope, v_plan_next_due và v_remediation_worklist không làm tăng số bảng vật lý.",
    "Ở mức phân tích, mỗi bảng cốt lõi": "Ở mức phân tích, mỗi bảng cốt lõi tương ứng một thực thể hoặc quan hệ có quy tắc riêng. Lớp miền không sao chép máy móc cấu trúc SQL: phê duyệt phạm vi, khởi tạo lần quét, xác nhận mức độ và kết luận xác minh thuộc dịch vụ nghiệp vụ; tầng truy cập dữ liệu chịu trách nhiệm lưu trữ và giao dịch. Bảng dưới nêu các lớp và quan hệ cần thể hiện khi dựng UML chi tiết.",
    "Hồ sơ UI có 32 màn hình": "Hồ sơ giao diện có 32 màn hình và 57 thao tác trong đặc tả OpenAPI. Điều hướng theo tài sản, lịch và lần quét, lô nhập, phát hiện, phiếu khắc phục, xác minh và báo cáo. Các màn hình rủi ro và hộp thông báo thuộc phần mở rộng khi áp dụng mô hình lõi 13 bảng. Giao diện hiển thị tách biệt Risk, Confidence và Final Severity; báo cáo gốc chỉ được mở qua quyền riêng và có ghi nhật ký kiểm toán.",
    "Đến thời điểm tổng hợp": "Đến thời điểm tổng hợp, dự án đã có hồ sơ ca sử dụng, sơ đồ hoạt động, OpenAPI, thiết kế giao diện và ERD. Giao diện web và API đang ở mức nền phát triển; tiến trình nền điều phối ZAP đã có mã và kiểm thử đơn vị cho lịch, kế hoạch quét, kiểm tra báo cáo và tham chiếu bí mật. Chưa có bằng chứng hoàn tất bộ phân tích báo cáo F-03, giao diện nghiệp vụ, tích hợp đầu cuối hoặc chuyển đổi cơ sở dữ liệu 13 bảng. Tiến trình nền mặc định tắt; chỉ nên bật sau khi thử trên mục tiêu được phép ở môi trường thử nghiệm.",
    "Các kiểm tra tiếp theo": "Kiểm thử tiếp theo cần bao phủ quyền theo tài sản, phạm vi URL và khung giờ, hoàn tác khi báo cáo lỗi, chống tạo lịch hoặc lô nhập trùng, xử lý quá hạn và mất quyền giữ công việc, tính toàn vẹn tệp báo cáo, quyết định Final Severity và xác minh đóng hoặc mở lại phát hiện. Kết quả phải ghi rõ môi trường và giới hạn kiểm thử, không coi hồ sơ thiết kế là bằng chứng nghiệm thu vận hành.",
}

for p in doc.paragraphs:
    if not paragraph_in_chapters(p) or not p.text.strip(): continue
    text = p.text
    for start,rewrite in specific.items():
        if text.startswith(start): text = rewrite; break
    text = vi(text)
    for start,rewrite in final_rewrites.items():
        if text.startswith(start): text = rewrite; break
    text = text.replace("3.8. Sơ đồ lớp", "3.7. Sơ đồ lớp")
    text = text.replace("3.9. Thiết kế giao diện", "3.8. Thiết kế giao diện")
    text = text.replace("3.10. Tình trạng", "3.9. Tình trạng")
    text = text.replace("Bảng 3.7. Lớp miền", "Bảng 3.6. Lớp miền")
    text = text.replace("Bảng 3.8. Màn hình", "Bảng 3.7. Màn hình")
    text = text.replace("Danh mục 63 ca sử dụng từ hồ sơ thiết kế ban đầu và phạm vi mô hình đích", "Danh mục 63 ca sử dụng trong hồ sơ thiết kế")
    text = text.replace("Danh mục 16 sơ đồ hoạt động thiết kế ban đầu", "Danh mục 16 sơ đồ hoạt động")
    for number in (1,2,3):
        text = text.replace(f"Hình 3.{number}. ca sử dụng", f"Hình 3.{number}. Sơ đồ ca sử dụng")
    if text != p.text: p.text = text

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if p.text.strip():
                    new = vi(p.text)
                    if new != p.text: p.text = new
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(12)

for table in doc.tables:
    if table.rows[0].cells[0].text == "Thành phần" and len(table.columns) == 3:
        table.rows[1].cells[0].text = "Giao diện web"
    if table.rows[0].cells[0].text == "Thứ tự" and len(table.columns) == 3:
        sequence_rows = [
            ("Bộ lập lịch → tầng dữ liệu", "Chọn lịch đủ điều kiện, khóa kỳ quét, tạo hồ sơ lần quét và phiên thực thi."),
            ("Tiến trình nền → bộ tạo kế hoạch", "Đọc bản chụp phạm vi và hồ sơ quét, tạo kế hoạch YAML; lấy bí mật qua tham chiếu."),
            ("Tiến trình nền → bộ thực thi ZAP", "Khởi chạy vùng chứa biệt lập, theo dõi thời gian tối đa và tín hiệu duy trì."),
            ("Bộ thực thi → bộ kiểm tra báo cáo", "Trả báo cáo định dạng Traditional; kiểm nguồn, tên miền, kích thước và SHA-256."),
            ("Kho tệp → tầng dữ liệu", "Lưu tệp báo cáo bất biến; cập nhật lần quét Completed và tạo lô nhập Received cùng giao dịch."),
            ("Bộ phân tích F-03 → phát hiện", "Đọc lô nhập, chuẩn hóa và chống trùng; phần triển khai hoàn chỉnh còn cần hoàn tất."),
        ]
        for i,(component,result) in enumerate(sequence_rows,1):
            table.rows[i].cells[1].text = component
            table.rows[i].cells[2].text = result

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(12)

# Nội dung các chương sau giữ nguyên; yêu cầu cỡ chữ bảng áp dụng cả bảng mẫu ở Chương V.
for p in doc.paragraphs:
    if p.text.startswith("    3.5–3.7."):
        p.text = "    3.5–3.6. ERD và danh sách bảng dữ liệu"
    elif p.text.startswith("    3.8–3.10."):
        p.text = "    3.7–3.9. Lớp miền, giao diện và kiểm chứng"
    elif p.text.startswith("    3.1–3.4."):
        p.text = "    3.1–3.4. Ca sử dụng, hoạt động, trạng thái và kiến trúc"

# Làm mới danh sách hình/bảng ở đầu tài liệu sau khi xóa và đánh lại số.
def h1(prefix):
    return next(p for p in doc.paragraphs if p.style.name == "Heading 1" and p.text.startswith(prefix))
heading_list = h1("DANH SÁCH HÌNH, BẢNG")
heading_abbr = h1("DANH MỤC TỪ VIẾT TẮT")
elems = list(body)
for element in elems[elems.index(heading_list._p)+1:elems.index(heading_abbr._p)]:
    body.remove(element)
figures=[]; tables=[]
for p in doc.paragraphs:
    if paragraph_in_chapters(p):
        if p.text.startswith("Hình "): figures.append(p.text)
        if p.text.startswith("Bảng "): tables.append(p.text)
for item in figures+tables:
    p=doc.add_paragraph(item)
    heading_abbr._p.addprevious(p._p)

doc.save(OUT)
print(OUT)
print("UC columns",len(uc.columns),"AD columns",len(ad.columns),"figures",len(figures),"table labels",len(tables))
