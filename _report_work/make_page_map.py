import json
import re
import sys
from pathlib import Path

from docx import Document
from pypdf import PdfReader

root = Path(r"C:\Users\Admin\Documents\ChatGPT\HTQLHDK")
docx = root / "outputs" / "BaoCaoPT_TKHTTT_Nhom20_HoanThien_6Chuong.docx"
pdf = root / "_report_work" / (sys.argv[1] if len(sys.argv) > 1 else "render_complete_1") / "BaoCaoPT_TKHTTT_Nhom20_HoanThien_6Chuong.pdf"
doc = Document(docx)
reader = PdfReader(pdf)
start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "MỤC LỤC")
end = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "DANH MỤC TỪ VIẾT TẮT")
keys = list(dict.fromkeys(
    p.text.split("\t")[0].strip()
    for p in doc.paragraphs[start + 1:end]
    if p.text.strip()
))
keys.extend(["DANH MỤC TỪ VIẾT TẮT"])

def normalize(value):
    return re.sub(r"\s+", " ", value).strip()

pages = {}
texts = [normalize(p.extract_text()) for p in reader.pages]
for key in keys:
    hits = [i + 1 for i, text in enumerate(texts) if normalize(key) in text]
    pages[key] = max(hits) if hits else None

path = root / "_report_work" / "page_map.json"
path.write_text(json.dumps(pages, ensure_ascii=False, indent=2), encoding="utf-8")
print("keys", len(keys), "missing", [k for k, v in pages.items() if v is None])
print("chapter pages", [(k, v) for k, v in pages.items() if k.startswith("CHƯƠNG")])
