from pathlib import Path
import math, re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\Admin\Documents\ChatGPT\HTQLHDK")
SOURCE = ROOT / "outputs" / "BaoCaoPT&TKHTTT_Nhom20_28092026.docx"
OUT = ROOT / "outputs" / "BaoCaoPT_TKHTTT_Nhom20_28092026_ban_viet_lai.docx"
WORK = ROOT / "_report_work"
WORK.mkdir(exist_ok=True)

def draw_erd(path):
    W, H = 2300, 1500
    im = Image.new("RGB", (W, H), "white")
    dr = ImageDraw.Draw(im)
    font_path = r"C:\Windows\Fonts\arial.ttf"
    bold_path = r"C:\Windows\Fonts\arialbd.ttf"
    f = ImageFont.truetype(font_path, 30)
    fb = ImageFont.truetype(bold_path, 36)
    title = ImageFont.truetype(bold_path, 48)
    dr.text((80, 25), "ERD khái quát - mô hình đích 13 bảng", font=title, fill="#17324d")
    cols = [80, 625, 1170, 1715]
    rows = [145, 445, 745, 1045]
    bw, bh = 455, 205
    nodes = {
        "app_user": (0,0,"Định danh"), "user_role": (0,1,"Quyền theo tài sản"),
        "asset": (0,2,"Website"), "audit_log": (0,3,"Nhật ký bất biến"),
        "asset_scope_revision": (1,0,"Phiên bản phạm vi"),
        "asset_scope_rule": (1,1,"Luật include/exclude"),
        "asset_scan_window": (1,2,"Khung giờ quét"),
        "zap_scan_profile": (1,3,"Hồ sơ ZAP"),
        "scan_plan": (2,1,"Lịch / xác minh"), "scan_run": (2,2,"Một lần chạy"),
        "finding": (3,0,"Phát hiện"),
        "finding_occurrence": (3,2,"Lần xuất hiện"),
        "remediation_ticket": (3,3,"Phiếu khắc phục"),
    }
    centers = {n:(cols[c]+bw/2,rows[r]+bh/2) for n,(c,r,_) in nodes.items()}
    def edge(a,b,label=None):
        x1,y1=centers[a]; x2,y2=centers[b]
        ang=math.atan2(y2-y1,x2-x1)
        x1+=math.cos(ang)*235; y1+=math.sin(ang)*110
        x2-=math.cos(ang)*235; y2-=math.sin(ang)*110
        dr.line((x1,y1,x2,y2), fill="#7990a4", width=5)
        aa=math.atan2(y2-y1,x2-x1)
        for shift in (-0.55,0.55):
            ex=x2-22*math.cos(aa+shift); ey=y2-22*math.sin(aa+shift)
            dr.line((x2,y2,ex,ey), fill="#7990a4", width=5)
        if label:
            mx,my=(x1+x2)/2,(y1+y2)/2
            dr.rounded_rectangle((mx-70,my-19,mx+70,my+19),8,fill="white")
            dr.text((mx-59,my-16),label,font=f,fill="#35516b")
    for a,b in [
        ("app_user","user_role"),("app_user","asset"),("app_user","audit_log"),
        ("asset","asset_scope_revision"),("asset_scope_revision","asset_scope_rule"),
        ("asset_scope_revision","asset_scan_window"),("asset_scope_revision","scan_plan"),
        ("zap_scan_profile","scan_plan"),("scan_plan","scan_run"),
        ("scan_run","finding_occurrence"),("finding","finding_occurrence"),
        ("finding","remediation_ticket")]: edge(a,b)
    for name,(c,r,desc) in nodes.items():
        x,y=cols[c],rows[r]
        dr.rounded_rectangle((x,y,x+bw,y+bh),22,fill="#edf4f8",outline="#315f7a",width=5)
        dr.rounded_rectangle((x,y,x+bw,y+72),22,fill="#d5e6f0")
        dr.text((x+18,y+17),name,font=fb,fill="#14334a")
        dr.text((x+18,y+105),desc,font=f,fill="#223e50")
    dr.text((90,1360),"Quan hệ xác minh: remediation_ticket -> scan_plan (plan_kind = Verification)",font=f,fill="#35516b")
    dr.text((90,1410),"Mỗi bảng có khóa chính id; khóa ngoại và ràng buộc chi tiết trình bày tại Bảng 3.5.",font=f,fill="#35516b")
    im.save(path)

ERD = WORK / "erd_13_bang_bao_cao.png"
draw_erd(ERD)

doc = Document(SOURCE)
body = doc._element.body
for child in list(body):
    if child.tag != qn("w:sectPr"): body.remove(child)
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.54)
sec.header_distance = sec.footer_distance = Cm(1.27)
sec.different_first_page_header_footer = True

normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Times New Roman", Pt(13)
normal.font.color.rgb = RGBColor(0,0,0)
normal.paragraph_format.line_spacing = 1.25
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name,size in [("Heading 1",16),("Heading 2",14)]:
    s=doc.styles[name]
    s.font.name,s.font.size,s.font.bold="Times New Roman",Pt(size),True
    s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.keep_with_next=True
    s.paragraph_format.space_before=Pt(14 if name=="Heading 1" else 9)
    s.paragraph_format.space_after=Pt(6)

def para(text="", bold=False, align=None, indent=True, style=None):
    q=doc.add_paragraph(style=style)
    if indent and style is None: q.paragraph_format.first_line_indent=Cm(.7)
    if align is not None: q.alignment=align
    rr=q.add_run(text); rr.bold=bold
    return q
def h1(s): doc.add_heading(s,level=1)
def h2(s): doc.add_heading(s,level=2)
def point(s):
    q=para("• "+s,indent=False)
    q.paragraph_format.left_indent=Cm(.5)
    q.paragraph_format.first_line_indent=Cm(-.35)
    return q
def page(): doc.add_page_break()

TABLES=[]; FIGS=[]
def tab(label, headers, rows, widths=None, fs=10):
    TABLES.append(label)
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    if widths:
        for i,width in enumerate(widths): t.columns[i].width=Cm(width)
    br=OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e=OxmlElement("w:"+edge); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"4"); e.set(qn("w:color"),"666666"); br.append(e)
    t._tbl.tblPr.append(br)
    for i,v in enumerate(headers): t.rows[0].cells[i].text=str(v)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v)
    for j,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        no=OxmlElement("w:cantSplit"); trPr.append(no)
        if j==0:
            th=OxmlElement("w:tblHeader"); th.set(qn("w:val"),"true"); trPr.append(th)
        for i,cell in enumerate(row.cells):
            if widths: cell.width=Cm(widths[i])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for q in cell.paragraphs:
                q.paragraph_format.space_after=Pt(1)
                q.paragraph_format.line_spacing=1.05
                for rr in q.runs:
                    rr.font.name="Times New Roman"; rr.font.size=Pt(fs)
                    if j==0: rr.bold=True
    para(label,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False,style="Caption1")
    return t
def fig(path,label,width=15.5):
    FIGS.append(label)
    q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.keep_with_next=True
    q.add_run().add_picture(str(path),width=Cm(width))
    para(label,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False,style="Caption1")

# Bìa
for s in ["BỘ KHOA HỌC VÀ CÔNG NGHỆ", "HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG", "-----------------------------"]:
    q=para(s,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
    for rr in q.runs: rr.font.size=Pt(15)
for _ in range(4): para("",indent=False)
q=para("BÁO CÁO\nĐỒ ÁN MÔN HỌC",bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
for rr in q.runs: rr.font.size=Pt(28)
para("MÔN HỌC: PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG THÔNG TIN",bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
para("ĐỀ TÀI: HỆ THỐNG QUẢN LÝ QUÉT LỖ HỔNG WEBSITE ĐỊNH KỲ",bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
para("Nhóm 20 · Báo cáo Chương I–III · Cập nhật 28/09/2026",align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
for _ in range(3): para("",indent=False)
para("Giảng viên hướng dẫn: [Điền họ tên]",indent=False)
para("Thành viên nhóm 20: [Điền họ tên và MSSV]",indent=False)
para("Lớp: [Điền lớp sinh viên]",indent=False)
for _ in range(3): para("",indent=False)
para("TP. Hồ Chí Minh, tháng 09/2026",align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
page()

h1("MỤC LỤC")
for ch,subs in [
    ("CHƯƠNG I. TỔNG QUAN",["Bối cảnh và mục tiêu", "Phạm vi và nền tảng kỹ thuật"]),
    ("CHƯƠNG II. PHÂN TÍCH NỘI DUNG, YÊU CẦU",["Tác nhân và quy trình nghiệp vụ", "Yêu cầu chức năng, dữ liệu và chất lượng"]),
    ("CHƯƠNG III. PHÂN TÍCH THIẾT KẾ HỆ THỐNG",["Use case và luồng hoạt động", "Trạng thái, kiến trúc và giao diện", "ERD và danh sách 13 bảng"]),
]:
    para(ch,bold=True,indent=False)
    for sub in subs: para("    "+sub,indent=False)
para("TÀI LIỆU THAM KHẢO",bold=True,indent=False)
h1("DANH SÁCH HÌNH, BẢNG")
fig_list=para("",indent=False)
table_list=para("",indent=False)
h1("DANH MỤC TỪ VIẾT TẮT")
tab("Bảng 0.1. Danh mục từ viết tắt",["Ký hiệu","Giải thích"],[
    ("HTQLHDK","Hệ thống quản lý quét lỗ hổng website định kỳ"),
    ("ATTT","An toàn thông tin"),("OWASP ZAP","OWASP Zed Attack Proxy"),
    ("UC / AD","Use case / activity diagram"),("ERD","Sơ đồ thực thể–quan hệ"),
    ("BCNF","Dạng chuẩn Boyce–Codd"),("API","Giao diện lập trình ứng dụng"),
    ("Worker","Tiến trình xử lý nền"),("FK / PK","Khóa ngoại / khóa chính"),
], [3.2,12.6])
page()

h1("CHƯƠNG I. TỔNG QUAN")
h2("1.1. Bối cảnh và lý do chọn đề tài")
para("Một website có thể sinh nhiều phát hiện bảo mật qua các kỳ quét. Nếu chỉ lưu tệp kết quả, đơn vị khó trả lời phát hiện thuộc tài sản nào, phạm vi quét có được chấp thuận hay không, ai chịu trách nhiệm và việc khắc phục đã được xác minh chưa. HTQLHDK giải quyết bài toán quản lý vòng đời đó, từ khai báo tài sản đến theo dõi và đóng phát hiện [1].")
para("Phần cập nhật ngày 28/09/2026 đưa thêm điều phối OWASP ZAP tự động qua Worker và đề xuất mô hình dữ liệu đích 13 bảng. Hai phần này ở trạng thái khác nhau: Worker và migration V002 đã có mã nguồn, còn mô hình 13 bảng là thiết kế tối giản chưa được chuyển thành SQL hiện hành. Báo cáo trình bày cả yêu cầu đích và hiện trạng một cách tách bạch [2], [3].")
h2("1.2. Mục tiêu và kết quả kỳ vọng")
for s in [
    "Quản lý tài sản website cùng phiên bản phạm vi, luật include/exclude và khung giờ quét được phê duyệt.",
    "Lập lịch định kỳ hoặc quét xác minh, điều phối ZAP theo hồ sơ cấu hình đã duyệt và giám sát từng lần chạy.",
    "Nhận báo cáo nguồn có checksum, kiểm định dạng/phạm vi, chuẩn hóa phát hiện và chống nhập trùng.",
    "Tách mức độ nguồn của ZAP khỏi Final Severity do ATTT quyết định; giao việc, theo dõi và xác minh khắc phục.",
    "Bảo vệ báo cáo lỗ hổng bằng phân quyền theo vai trò, tài sản và nhiệm vụ, che dữ liệu nhạy cảm và ghi audit.",
]: point(s)
h2("1.3. Phạm vi và giới hạn")
para("Bốn chức năng cốt lõi là F-01 quản lý tài sản/phạm vi, F-02 lịch và lần quét, F-03 tiếp nhận kết quả, F-04 khắc phục/xác minh. OWASP ZAP là scanner bên ngoài duy nhất. Hệ thống không được quét một URL chỉ vì URL đó tồn tại trong dữ liệu; tài sản, phiên bản phạm vi, thời điểm và profile phải đồng thời đủ điều kiện [1], [3].")
para("Bộ use case gốc có 63 mục, gồm các quy trình mở rộng như chấp nhận rủi ro, xuất báo cáo và thông báo. Mô hình đích 13 bảng tập trung bốn chức năng cốt lõi, bỏ lưu trữ các quy trình mở rộng này; một số nhu cầu đọc/xuất có thể tạo bằng truy vấn hoặc xử lý theo yêu cầu. Vì vậy các sơ đồ gốc được dùng để truy vết yêu cầu, còn ERD 13 bảng đại diện cho phạm vi dữ liệu đích [3], [4].")
h2("1.4. Cơ sở lý thuyết")
para("Một phát hiện bảo mật trải qua phát hiện, phân tích, phân công, khắc phục và xác minh. Dữ liệu ZAP như Risk và Confidence phản ánh đánh giá của công cụ; Final Severity là quyết định nghiệp vụ có lý do và người xác nhận. Fingerprint kết hợp tài sản và đặc điểm lỗi để chống trùng qua các lần quét. Việc phát hiện không còn xuất hiện chỉ đủ làm bằng chứng khi vùng bao phủ quét xác minh đã được xác nhận [1], [3].")
para("Thiết kế cơ sở dữ liệu dùng khóa, ràng buộc duy nhất và quan hệ tách biệt cho dữ liệu nhiều giá trị. Mô hình đích hướng đến BCNF cho phần lớn bảng, không lưu thông tin tính được như kỳ quét kế tiếp, first/last seen hay số liệu dashboard. Audit append-only bảo lưu lịch sử thay đổi nhạy cảm [3].")
h2("1.5. Nền tảng kỹ thuật và hiện trạng")
tab("Bảng 1.1. Thành phần kỹ thuật và mức hoàn thiện",["Thành phần","Lựa chọn","Hiện trạng"],[
    ("Web UI","React 19.2, TypeScript, Vite","Có nền ứng dụng; 32 màn hình đang ở mức hồ sơ thiết kế."),
    ("API","ASP.NET Core/.NET 10; REST/OpenAPI 3.1","Có nền API; 57 operation ở mức hợp đồng thiết kế."),
    ("Worker",".NET 10, Npgsql, Docker ZAP runner","Có mã điều phối ZAP; mặc định tắt khi chưa cấu hình."),
    ("Cơ sở dữ liệu","PostgreSQL 17","V001: 27 bảng; V002 bổ sung 3 bảng ZAP."),
    ("Mô hình đích","13 bảng BCNF tối giản","Tài liệu thiết kế; chưa có migration chuyển đổi."),
], [3.0,5.0,7.8],10)
fig(ROOT/"architecture/rendered/02_container_components.png","Hình 1.1. Kiến trúc thành phần nền của HTQLHDK (nguồn: [5]).",15.5)

page()
h1("CHƯƠNG II. PHÂN TÍCH NỘI DUNG, YÊU CẦU")
h2("2.1. Tác nhân, quyền và trách nhiệm")
tab("Bảng 2.1. Tác nhân nghiệp vụ và tác nhân kỹ thuật",["Tác nhân","Vai trò trong quy trình","Giới hạn quyết định"],[
    ("Chủ tài sản","Khai báo, xác nhận sở hữu/phạm vi; phản hồi tác động","Không tự phê duyệt phạm vi quét."),
    ("Chuyên viên ATTT","Phê duyệt scope, đánh giá mức độ, giao việc, kết luận xác minh","Không bỏ qua điều kiện an toàn và audit."),
    ("Người vận hành quét","Tạo/theo dõi lịch, xử lý lỗi kỹ thuật và báo cáo","Không tự quyết định Final Severity."),
    ("Người khắc phục","Nhận phiếu, cập nhật biện pháp và bằng chứng","Không tự đóng finding."),
    ("Quản trị viên / kiểm toán","Cấp/thu hồi quyền; xem lịch sử được phép","ADMIN không mặc nhiên đọc raw report."),
    ("Worker / OWASP ZAP","Điều phối và thực hiện quét kỹ thuật","Không phê duyệt nghiệp vụ hoặc tự đóng finding."),
], [3.0,6.5,6.3],9.7)
h2("2.2. Quy trình nghiệp vụ tổng thể")
para("Luồng chính bắt đầu khi chủ tài sản khai báo website và phạm vi; ATTT phê duyệt; người vận hành hoặc Worker tạo lần quét; ZAP xuất báo cáo; hệ thống kiểm tra và ghi phát hiện; ATTT rà soát; người khắc phục xử lý; cuối cùng quét xác minh dẫn tới đóng hoặc mở lại. Mọi nhánh từ chối do sai quyền, trạng thái hoặc phạm vi phải dừng trước khi ghi dữ liệu nghiệp vụ và có audit phù hợp [1], [2].")
fig(ROOT/"so_do_cap_nhat/rendered/02_activity_quy_trinh_tong_the.png","Hình 2.1. Luồng hoạt động tổng thể tích hợp ZAP tự động (nguồn: [2]).",13.2)
h2("2.3. Quy trình F-01: tài sản và phạm vi")
tab("Bảng 2.2. Luồng F-01",["Bước","Hành động","Điều kiện/kết quả"],[
    ("1","Chủ tài sản hoặc ATTT khai báo asset","Mã duy nhất; URL, chủ sở hữu, mức quan trọng đầy đủ."),
    ("2","Tạo scope revision, rule và scan window","Bản nháp chưa tác động lịch đang chạy."),
    ("3","Chủ tài sản xác nhận","Xác nhận sở hữu, mức quan trọng và phạm vi."),
    ("4","ATTT phê duyệt/từ chối","Approved bất biến và có hiệu lực; từ chối giữ lịch sử và lý do."),
    ("5","Tạm dừng/ngừng tài sản khi cần","Không dispatch quét mới khi tài sản không Active."),
], [1.2,5.9,8.7],10)
para("Trùng domain/URL với một tài sản khác là tín hiệu để rà soát, không tự động từ chối mọi trường hợp. Một thay đổi phạm vi tạo revision mới; lịch đang có chỉ dùng phiên bản đã được duyệt và còn hiệu lực [6].")
h2("2.4. Quy trình F-02: lịch và điều phối quét")
tab("Bảng 2.3. Luồng F-02",["Bước","Hành động","Guard và nhánh ngoại lệ"],[
    ("1","Lập hoặc sửa lịch/profile","Tài sản Active, scope Approved, thời gian và quyền hợp lệ."),
    ("2","Đến hạn, tạo occurrence/execution","Khóa kỳ quét chống trùng; lịch manual không do Worker dispatch."),
    ("3","Worker dựng Automation Framework plan","Lấy URL/rule từ scope, secret chỉ qua tham chiếu; active scan mặc định tắt."),
    ("4","Chạy container ZAP biệt lập","Giới hạn tài nguyên, timeout, lease và heartbeat."),
    ("5","Thu báo cáo, ghi kết quả hoặc lỗi","Không tự retry lần quét đã có thể gửi request; lỗi được che."),
], [1.2,5.9,8.7],9.7)
para("V002 tách scan_occurrence (trạng thái nghiệp vụ) và zap_execution (trạng thái kỹ thuật). Sau khi báo cáo hợp lệ, Worker lưu artifact và tạo import_batch Received trong cùng giao dịch; parser F-03 tiếp tục xử lý. Kế hoạch 13 bảng sau này hợp nhất các quan hệ 1:1 này thành scan_run, nhưng hiện chưa triển khai [2], [3].")
h2("2.5. Quy trình F-03: báo cáo và phát hiện")
tab("Bảng 2.4. Luồng F-03",["Bước","Nội dung","Kết quả/ngoại lệ"],[
    ("1","Nhận tệp hoặc artifact từ Worker","Gắn đúng asset và lần quét; kiểm quyền với nhánh manual."),
    ("2","Kiểm loại, kích thước, host, checksum","Từ chối báo cáo sai cấu trúc/phạm vi; không tạo finding."),
    ("3","Lưu báo cáo gốc bất biến","Metadata và SHA-256 đi cùng báo cáo; không log payload thô."),
    ("4","Chuẩn hóa và chống trùng","Fingerprint theo tài sản; ghi occurrence cho từng lần chạy."),
    ("5","ATTT rà soát","Phân biệt Source Risk, Confidence và Final Severity; xác định false positive."),
], [1.2,5.9,8.7],9.7)
para("Hợp đồng baseline 27 bảng cho phép lô nhập một phần và lưu import_error; mô hình đích 13 bảng chủ ý xử lý nguyên tử: bất kỳ lỗi cấu trúc/phạm vi khiến toàn bộ báo cáo bị từ chối hoặc rollback. Đây là thay đổi hành vi cần cập nhật parser, API và kiểm thử trước khi áp dụng [3], [6].")
h2("2.6. Quy trình F-04: khắc phục và xác minh")
tab("Bảng 2.5. Luồng F-04",["Bước","Vai trò/hành động","Điều kiện kết thúc"],[
    ("1","ATTT tạo và giao phiếu","Finding đủ điều kiện; người phụ trách và hạn hợp lệ."),
    ("2","Người khắc phục tiếp nhận","Cập nhật nguyên nhân, biện pháp, tiến độ, bằng chứng."),
    ("3","Chuyển chờ xác minh","Thiếu biện pháp/bằng chứng thì giữ In Progress."),
    ("4","Lập quét xác minh","Kết quả gắn đúng finding, tài sản và phạm vi."),
    ("5","ATTT kết luận","Đóng khi coverage và kết quả hợp lệ; còn lỗi/mất điều kiện thì Reopened."),
], [1.2,6.0,8.6],9.7)
h2("2.7. Yêu cầu chức năng và dữ liệu vào/ra")
tab("Bảng 2.6. Ma trận yêu cầu chức năng",["Nhóm","Đầu vào","Đầu ra","Use case tiêu biểu"],[
    ("F-01","Asset, URL, owner, rule, window","Asset/scope Approved","UC-A01–A07"),
    ("F-02","Plan, RRULE, profile, scope","Scan run, trạng thái/lỗi","UC-S01–S08"),
    ("F-03","Traditional JSON/XML, metadata","Finding, occurrence, lỗi nhập","UC-I01–I12; T01–T06"),
    ("F-04","Finding, assignee, evidence","Ticket, kết quả xác minh","UC-R01–R13"),
    ("Hỗ trợ","Role, asset, action","Dữ liệu đã che, audit","UC-C01–C10"),
], [2.1,4.3,4.9,4.5],9.4)
h2("2.8. Quy tắc nghiệp vụ và yêu cầu chất lượng")
for s in [
    "Chỉ quét tài sản Active trong scope Approved và scan window hợp lệ; luật Exclude được ưu tiên khi URL khớp cả Include và Exclude.",
    "Profile ZAP có phiên bản; bản Active/Retired bất biến. Active scan chỉ bật sau phê duyệt mục tiêu và cấu hình.",
    "Tách tài khoản dịch vụ khỏi tài khoản con người; tài khoản dịch vụ không đăng nhập giao diện hoặc phê duyệt.",
    "Báo cáo nguồn và metadata checksum bất biến; dữ liệu Attack/Evidence, token và URL nhạy cảm không xuất hiện trong log/audit thô.",
    "API kiểm lại quyền theo action, tài sản, nhiệm vụ và trạng thái; UI chỉ giúp định hướng, không phải ranh giới bảo mật.",
    "Các thao tác ghi cần idempotency, transaction và audit; mất lease không được dẫn tới quét lại âm thầm.",
    "Giao diện phải có trạng thái chờ, lỗi, trùng, từ chối và hỗ trợ bàn phím; thời gian hiển thị kèm múi giờ.",
]: point(s)
para("Mục tiêu định lượng như SLA, thời hạn lưu báo cáo, giới hạn upload và RPO/RTO chưa được phê duyệt trong hồ sơ; báo cáo không tự đặt con số. Kiểm thử chấp nhận phải đi qua đường thành công và các đường từ chối, rollback, kiểm quyền và chạy đồng thời [5], [7].")

page()
h1("CHƯƠNG III. PHÂN TÍCH THIẾT KẾ HỆ THỐNG")
h2("3.1. Ranh giới hệ thống và use case")
para("Sơ đồ use case phân vai người dùng theo trách nhiệm. Những UC kiểm định dạng, chống trùng, masking và audit là hành vi nội bộ được use case cha gọi, không phải nút thao tác tự do. Bộ sơ đồ gốc gồm 63 UC; nhóm RA về chấp nhận rủi ro và một số chức năng báo cáo/thông báo là phạm vi mở rộng so với ERD đích 13 bảng [4], [6].")
fig(ROOT/"visual_paradigm_uml_md/rendered/18_use_case_tai_san_lich_quet.png","Hình 3.1. Use case tài sản và lịch quét baseline (nguồn: [6]).",8.2)
fig(ROOT/"visual_paradigm_uml_md/rendered/19_use_case_nhap_bao_cao_zap.png","Hình 3.2. Use case nhập báo cáo ZAP baseline (nguồn: [6]).",11.3)
fig(ROOT/"visual_paradigm_uml_md/rendered/21_use_case_khac_phuc_xac_minh.png","Hình 3.3. Use case khắc phục và xác minh baseline (nguồn: [6]).",9.2)
para("Trong mô hình đích, use case quản lý tài sản, lập kế hoạch, nhận báo cáo và khắc phục được giữ. Chức năng chấp nhận rủi ro, SLA policy, notification inbox và link chia sẻ báo cáo không có bảng riêng; nếu được yêu cầu trở lại, phải bổ sung mô hình quan hệ và luồng phê duyệt tương ứng [3].")

UC=[]
for line in (ROOT/"tai_lieu_use_case/README.md").read_text(encoding="utf-8").splitlines():
    if line.startswith("| UC-"):
        c=[x.strip() for x in line.strip().strip("|").split("|")]
        if len(c)>=3:
            code,name,actor=c[:3]
            actor=actor.replace("Không có actor trực tiếp; hành vi nội bộ do UC cha gọi","Nội bộ")
            actor=actor.replace("Chủ sở hữu tài sản","Chủ tài sản").replace("Chuyên viên ATTT","ATTT")
            actor=actor.replace("Nhóm vận hành quét","Vận hành").replace("Nhóm khắc phục kỹ thuật","Khắc phục")
            actor=actor.replace("Người phê duyệt rủi ro","Duyệt rủi ro").replace("Kiểm toán viên","Kiểm toán")
            if len(actor)>65: actor=actor[:62]+"…"
            scope="Mở rộng" if code.startswith("UC-RA") else ("Hỗ trợ" if code.startswith("UC-C") else "Cốt lõi")
            UC.append((code,name,actor,scope))
tab("Bảng 3.1. Danh mục 63 use case từ hồ sơ baseline và phạm vi mô hình đích",["Mã","Use case","Tác nhân","Phạm vi"],UC,[2.2,6.2,5.3,2.1],8.4)

h2("3.2. Sơ đồ hoạt động và các nhánh quyết định")
para("Sơ đồ hoạt động ở Hình 2.1 bao quát luồng chính. Các sơ đồ chi tiết AD-02, AD-03, AD-04 và AD-06 chỉ rõ trách nhiệm từng swimlane, guard và nhánh lỗi. Khi bổ sung tự động hóa, nhánh quét ngoài của baseline được mở rộng bằng Worker/ZAP; quyết định nghiệp vụ của chủ tài sản và ATTT giữ nguyên [2], [6].")
AD=[]
for line in (ROOT/"tai_lieu_activity/README.md").read_text(encoding="utf-8").splitlines():
    if line.startswith("| AD-"):
        c=[x.strip() for x in line.strip().strip("|").split("|")]
        if len(c)>=2: AD.append((c[0],c[1],"Mở rộng" if c[0] in {"AD-08","AD-10","AD-13"} else "Cốt lõi/hỗ trợ"))
tab("Bảng 3.2. Danh mục 16 sơ đồ hoạt động baseline",["Mã","Luồng hoạt động","Phạm vi"],AD,[2.0,10.9,2.9],9.1)
para("Điểm quyết định quan trọng nhất của AD-02 là quyền khai báo, tính hợp lệ của URL và sự xác nhận/phê duyệt. AD-03 kiểm tài sản Active, scope Approved và xung đột lịch. AD-04 kiểm quyền, khóa lô, định dạng và URL thuộc phạm vi trước khi tạo finding. AD-06 chỉ cho đóng khi kết quả xác minh đúng tài sản/finding và đủ vùng bao phủ; nếu kết quả không hợp lệ hoặc fingerprint còn xuất hiện, phát hiện quay về Reopened [6].")
fig(ROOT/"visual_paradigm_uml_md/rendered/03_activity_khai_bao_phe_duyet_tai_san.png","Hình 3.4. Activity khai báo và phê duyệt tài sản (nguồn: [6]).",10.8)
fig(ROOT/"visual_paradigm_uml_md/rendered/07_activity_khac_phuc_va_xac_minh.png","Hình 3.5. Activity khắc phục và xác minh (nguồn: [6]).",10.7)

h2("3.3. Sơ đồ trạng thái và trình tự xử lý tự động")
para("Ở schema hiện hành, Worker chuyển zap_execution qua Queued → Preparing Plan → Starting Runner → Running → Collecting Report → Report Received; Failed, Timed Out và Cancelled là nhánh kết thúc lỗi. scan_occurrence vẫn là trạng thái nghiệp vụ; Worker chỉ tạo import_batch Received và không tự đóng finding [2].")
fig(ROOT/"so_do_cap_nhat/rendered/03_state_zap_execution.png","Hình 3.6. Trạng thái phiên ZAP trong V002 (nguồn: [2]).",15.5)
tab("Bảng 3.3. Trình tự các thành phần trong một lần quét tự động",["Thứ tự","Thành phần","Trao đổi/kết quả"],[
    ("1","Scheduler → Repository","Chọn lịch đủ điều kiện, khóa kỳ, tạo occurrence và execution."),
    ("2","Worker → Plan Renderer","Đọc snapshot scope/profile, dựng YAML; lấy secret qua tham chiếu."),
    ("3","Worker → ZAP Runner","Khởi chạy container biệt lập, theo dõi timeout và heartbeat."),
    ("4","Runner → Report Inspector","Trả Traditional report; kiểm nguồn, host, kích thước, checksum."),
    ("5","Artifact Store → Repository","Lưu artifact bất biến; cập nhật occurrence Completed và tạo batch Received cùng giao dịch."),
    ("6","Parser F-03 → Finding","Đọc batch, chuẩn hóa và chống trùng; hiện còn là phần triển khai tiếp."),
], [1.5,4.7,9.6],9.4)
h2("3.4. Thiết kế kiến trúc và luồng dữ liệu")
para("Web UI gọi API, API xác thực và kiểm quyền theo quan hệ vai trò–tài sản–nhiệm vụ. PostgreSQL giữ trạng thái nghiệp vụ, quyền và audit. Worker riêng điều phối ZAP; scanner không truy cập database nghiệp vụ. API không giữ Docker socket hoặc ZAP API key. Production cần runner quyền tối thiểu, object storage riêng tư và cơ chế quản lý khóa/secret được phê duyệt [2], [5].")
fig(ROOT/"so_do_cap_nhat/rendered/04_dfd_level_1.png","Hình 3.7. DFD mức 1 với bốn tiến trình nghiệp vụ (nguồn: [2]).",15.5)
h2("3.5. ERD mô hình đích 13 bảng")
para("ERD dưới đây thể hiện quan hệ cốt lõi của mô hình đích. app_user cấp quyền qua user_role và đóng vai trò owner/ATTT/operator/assignee; asset có nhiều phiên bản scope; scope liên kết luật và khung giờ; plan kết hợp scope với profile; run ghi một lần quét; finding_occurrence nối run với finding; ticket theo dõi khắc phục; audit_log lưu dấu vết. Quan hệ xác minh dùng scan_plan loại Verification tham chiếu ticket [3].")
fig(ERD,"Hình 3.8. ERD khái quát 13 bảng theo thiết kế đích (biên soạn từ [3]).",15.7)
tab("Bảng 3.4. Bội số và khóa ngoại chính trong ERD",["Quan hệ","Bội số","Khóa ngoại / ý nghĩa"],[
    ("app_user – user_role","1–N","user_role.user_id; quyền có thể giới hạn asset_id."),
    ("asset – scope_revision","1–N","asset_scope_revision.asset_id; phiên bản đã duyệt bất biến."),
    ("scope_revision – rule/window","1–N","scope_id ở hai bảng con; Exclude ưu tiên Include."),
    ("scope_revision/profile – plan","1–N","scan_plan.scope_id và zap_profile_id."),
    ("scan_plan – scan_run","1–N","scan_run.plan_id; Verification tối đa một run."),
    ("scan_run – finding","N–N","finding_occurrence là bảng nối hai FK."),
    ("finding – ticket","1–0..1","remediation_ticket.finding_id duy nhất trong mô hình tối giản."),
    ("ticket – Verification plan","1–N","scan_plan.remediation_ticket_id khi plan_kind=Verification."),
], [4.5,2.0,9.3],9.3)

h2("3.6. Danh sách bảng và thuộc tính dữ liệu chính")
para("Mọi bảng vật lý trong mô hình đích dùng id làm khóa chính. Bảng sau nêu thêm khóa ứng viên, các FK quan trọng và dữ liệu chính. Danh sách này mô tả thiết kế, không được hiểu là SQL hiện hành đã chuyển đổi [3].")
DATA=[
    ("app_user","username; account_kind; is_active","username duy nhất","Người dùng và service user; service không có mật khẩu giao diện."),
    ("user_role","user_id; role_code; asset_id; granted/revoked","user_id, asset_id","Quyền hiệu lực duy nhất theo user/role/asset."),
    ("asset","code; base_url; owner_id; security_officer_id; status","owner_id, security_officer_id","Website; code duy nhất, trạng thái Draft→Active/Suspended/Retired."),
    ("asset_scope_revision","asset_id; version_no; status; effective dates","asset_id","(asset_id,version_no) duy nhất; Approved bất biến."),
    ("asset_scope_rule","scope_id; effect; target_type/value; match_mode","scope_id","Mỗi include/exclude là một dòng, không nhét mảng/JSON."),
    ("asset_scan_window","scope_id; weekday; start/end; timezone","scope_id","Khung giờ hợp lệ gắn đúng phiên bản scope."),
    ("zap_scan_profile","code; version_no; spider/auth/report config","service_user_id, approved_by","Profile Active bất biến; secret chỉ lưu reference."),
    ("scan_plan","plan_kind; scope_id; profile_id; operator_id; RRULE","scope_id, zap_profile_id, remediation_ticket_id","Periodic hoặc Verification; cột theo loại có CHECK."),
    ("scan_run","plan_id; scheduled_for; scan/report status; SHA-256","plan_id, evaluated_by","Một lần chạy; report 1:1, kết quả xác minh nếu có."),
    ("finding","asset_id; fingerprint; final_severity; status","asset_id, assessed_by","(asset_id,fingerprint) duy nhất; mức cuối khác mức nguồn."),
    ("finding_occurrence","finding_id; scan_run_id; source risk/confidence","finding_id, scan_run_id","Bảng nối N–N; unique theo finding/run và record locator."),
    ("remediation_ticket","finding_id; assignee_id; due_at; progress","finding_id, assignee_id","Một phiếu hiện hành/finding; bằng chứng tối thiểu đã che."),
    ("audit_log","actor_id; action; object; outcome; masked before/after","actor_id","Append-only, correlation ID; không lưu payload thô."),
]
tab("Bảng 3.5. Từ điển tóm tắt 13 bảng vật lý của mô hình đích",["Bảng","Thuộc tính chính","FK tiêu biểu","Ràng buộc và vai trò"],DATA,[3.0,4.3,3.5,5.0],8.2)
para("Dữ liệu được suy ra thay vì lưu gồm scope hiện hành, kỳ quét kế tiếp, first_seen/last_seen, số phát hiện theo mức độ, worklist quá hạn và báo cáo tổng hợp. View được đề nghị như v_asset_current_scope, v_plan_next_due và v_remediation_worklist chỉ là đối tượng truy vấn, không làm tăng số bảng vật lý [3].")
h2("3.7. Đối chiếu ERD đích với SQL đang có")
tab("Bảng 3.6. Chuyển đổi từ schema hiện hành sang mô hình đích",["Hiện hành","Mô hình đích / tác động"],[
    ("V001: 27 bảng; V002: thêm 3 bảng","Tổng 30 bảng nghiệp vụ trước tinh gọn; chưa chuyển migration."),
    ("scan_schedule","Đổi thành scan_plan, gộp lịch định kỳ và yêu cầu xác minh."),
    ("scan_occurrence + zap_execution + import_batch + zap_artifact","Gộp thành scan_run; một lần quét có một báo cáo nguồn."),
    ("finding_review","Final Severity hiện hành ở finding; lịch sử ở audit_log."),
    ("ticket_event + evidence + verification_request/result","Tiến độ/bằng chứng ở ticket; xác minh bằng plan/run."),
    ("risk_acceptance, sla_rule, notification, report_export...","Bỏ khỏi lõi 13 bảng; cần thiết kế bổ sung nếu yêu cầu trở lại."),
], [6.3,9.5],9.0)
para("Chuyển đổi này không thể thực hiện bằng cách đổi tên bảng đơn thuần. Cần di trú dữ liệu lịch sử, cập nhật FK/API/Worker, thống nhất chính sách xử lý báo cáo nguyên tử và kiểm thử phân quyền, đối soát checksum, giao dịch cùng khả năng khôi phục lỗi. Mức 75/75 kiểm tra hiện có xác nhận bộ SQL 27 bảng ở phạm vi cấu trúc/ràng buộc đã thử, không xác nhận ERD 13 bảng đã chạy [3], [8].")
h2("3.8. Sơ đồ lớp miền nghiệp vụ")
para("Ở mức phân tích, mỗi bảng cốt lõi tương ứng một thực thể miền hoặc quan hệ có quy tắc riêng. Các lớp không được hiểu là một bản sao máy móc của cấu trúc SQL: hành vi phê duyệt phạm vi, khởi tạo lần quét, xác nhận mức độ và kết luận xác minh nằm trong dịch vụ nghiệp vụ, còn repository xử lý lưu trữ và giao dịch. Bảng dưới nêu quan hệ lớp cần thể hiện khi dựng UML chi tiết [3], [5].")
tab("Bảng 3.7. Lớp miền và hành vi chính",["Lớp / nhóm lớp","Hành vi chính","Quan hệ"],[
    ("Asset / ScopeRevision","Khai báo, xác nhận, phê duyệt phiên bản phạm vi","Asset 1–N ScopeRevision; revision 1–N rule/window."),
    ("ZapScanProfile / ScanPlan","Khóa profile đã duyệt; tạo lịch định kỳ hoặc xác minh","Profile và scope 1–N plan."),
    ("ScanRun","Ghi một lần chạy, checksum và kết quả xác minh","Plan 1–N run; Verification plan tối đa một run."),
    ("Finding / FindingOccurrence","Chống trùng, đánh giá mức độ, ghi lần xuất hiện","Run N–N Finding qua occurrence."),
    ("RemediationTicket / AuditLog","Theo dõi xử lý; lưu dấu vết bất biến","Finding 1–0..1 ticket; hành động → audit."),
], [4.2,5.7,5.9],9.1)
h2("3.9. Thiết kế giao diện và xử lý")
para("Hồ sơ UI có 32 màn hình và 57 operation OpenAPI. Điều hướng công việc gồm tài sản, lịch/lần quét, lô nhập, phát hiện, phiếu, xác minh, rủi ro và báo cáo. Khi chuyển về lõi 13 bảng, các màn hình rủi ro và inbox thông báo phải được loại khỏi bản lõi hoặc thiết kế lại theo chính sách đã duyệt. Giao diện luôn hiện Source Risk, Confidence và Final Severity thành trường riêng; raw report chỉ mở qua quyền riêng và audit [7].")
tab("Bảng 3.8. Màn hình ưu tiên và dữ liệu hiển thị",["Màn hình","Dữ liệu/hành động","Quy tắc kiểm soát"],[
    ("Tài sản/phạm vi","Asset, revision, rule, window; xác nhận/duyệt","Phiên Approved bất biến; actor đúng vai trò."),
    ("Lịch và lần quét","Plan, run, trạng thái ZAP; tạm dừng/xem lỗi","Chỉ hiển thị asset có quyền; không lộ secret."),
    ("Báo cáo/phát hiện","Checksum, Finding, Risk/Confidence/Final Severity","Raw/evidence che mặc định, xem riêng có audit."),
    ("Phiếu/xác minh","Assignee, hạn, biện pháp, coverage, quyết định","Đóng/mở lại chỉ sau kiểm tra ATTT."),
], [3.5,6.3,6.0],9.3)
h2("3.10. Tình trạng triển khai và kiểm chứng")
para("Đến thời điểm tổng hợp, tài liệu thiết kế use case, activity, OpenAPI, UI và ERD đã có. Backend/API/frontend đang ở mức nền phát triển; Worker điều phối ZAP đã có mã và kiểm thử đơn vị cho lịch, plan, report inspector, secret provider. Parser F-03 hoàn chỉnh, luồng UI nghiệp vụ, tích hợp end-to-end và chuyển migration 13 bảng chưa được chứng minh là hoàn tất. Worker mặc định tắt; cần thử trên staging được phép trước khi bật quét tự động [2], [3], [8].")
para("Các kiểm tra tiếp theo phải bao phủ quyền sai/đúng theo tài sản, phạm vi URL và khung giờ, rollback khi báo cáo lỗi, chống dispatch/import trùng, timeout/lease, độ toàn vẹn artifact, quyết định Final Severity, và xác minh đóng/mở lại finding. Kết quả kiểm tra phải ghi rõ môi trường và phạm vi, không dùng kết quả thiết kế làm bằng chứng nghiệm thu production.")

h1("TÀI LIỆU THAM KHẢO")
for s in [
    "[1] HTQLHDK, doi_tuong_sd_quy_trinh_hd_tich_hop_zap_tu_dong.md, yêu cầu nghiệp vụ cập nhật.",
    "[2] HTQLHDK, docs/ZAP_AUTOMATION_FRAMEWORK_CHANGES.md; so_do_cap_nhat/README.md.",
    "[3] HTQLHDK, database/08_thiet_ke_3nf_tinh_gon.md; database/09_giai_thich_chi_tiet_13_bang.md; database/README.md.",
    "[4] HTQLHDK, tai_lieu_use_case/README.md, danh mục 63 use case baseline.",
    "[5] HTQLHDK, architecture/README.md; architecture/01_kien_truc_logic_va_trien_khai.md.",
    "[6] HTQLHDK, tai_lieu_activity/README.md; visual_paradigm_uml_md/README.md.",
    "[7] HTQLHDK, ui_design/README.md; api_design/README.md.",
    "[8] HTQLHDK, docs/development/08_verification_report.md; database/validation/database_test_results.json.",
]: para(s,indent=False)

fig_list.text="\n".join(FIGS)
table_list.text="\n".join(TABLES)
sec.header.paragraphs[0].text="Báo cáo đồ án PT&TKHTTT · Nhóm 20 · HTQLHDK"
sec.header.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
for rr in sec.header.paragraphs[0].runs: rr.font.name="Times New Roman"; rr.font.size=Pt(9)
for q in sec.first_page_header.paragraphs: q.text=""
for q in sec.first_page_footer.paragraphs: q.text=""
f=sec.footer.paragraphs[0]; f.text=""; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); f._p.append(fld)
doc.core_properties.title="Báo cáo phân tích và thiết kế HTQLHDK - Chương I-III"
doc.core_properties.subject="Bản viết lại đầy đủ use case, activity, ERD và danh sách bảng"
doc.save(OUT)
print(OUT)
print("use_cases",len(UC),"activities",len(AD),"tables",len(TABLES),"figures",len(FIGS))
